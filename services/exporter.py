"""Write results to CSV or Excel, plus an optional easy-to-read TXT file."""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, List

COLUMNS = [
    "hr_name", "hr_name_source", "hr_email", "company_name", "company_domain", "company_website",
    "company_summary", "company_interest", "careers_page", "research_sources", "email_subject",
    "email_body", "llm_provider", "llm_model", "status", "notes",
]
BASE_NAME = "personalized_hr_emails"


def _csv_safe(value) -> str:
    """Stop web text like '=HYPERLINK(...)' from running as a spreadsheet formula."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def _write_with_fallback(path: Path, writer: Callable[[Path], None]) -> Path:
    """If the file is open in Excel (locked), save under a timestamped name instead."""
    try:
        writer(path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
        writer(alt)
        return alt


def write_csv(rows: List[Dict], path: Path) -> Path:
    def _w(p: Path) -> None:
        with p.open("w", encoding="utf-8-sig", newline="") as fh:  # utf-8-sig opens cleanly in Excel
            w = csv.DictWriter(fh, fieldnames=COLUMNS, extrasaction="ignore")
            w.writeheader()
            for row in rows:
                w.writerow({k: _csv_safe(row.get(k, "")) for k in COLUMNS})
    return _write_with_fallback(path, _w)


def write_xlsx(rows: List[Dict], path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    widths = {"hr_name": 18, "hr_name_source": 18, "hr_email": 30, "company_name": 22, "company_domain": 20,
              "company_website": 28, "company_summary": 45, "company_interest": 35, "careers_page": 30,
              "research_sources": 40, "email_subject": 40, "email_body": 80, "llm_provider": 12,
              "llm_model": 22, "status": 10, "notes": 40}
    status_fill = {"success": "C6EFCE", "partial": "FFEB9C", "failed": "FFC7CE"}

    def _w(p: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Emails"
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", start_color="1F3A5F")
        for r, row in enumerate(rows, start=2):
            for c, key in enumerate(COLUMNS, start=1):
                cell = ws.cell(row=r, column=c)
                cell.value = "" if row.get(key) is None else str(row.get(key))
                cell.data_type = "s"   # always plain text, never a formula
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            fill = status_fill.get(str(row.get("status", "")))
            if fill:
                ws.cell(row=r, column=COLUMNS.index("status") + 1).fill = PatternFill("solid", start_color=fill)
        for c, key in enumerate(COLUMNS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = widths.get(key, 20)
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(p)
    return _write_with_fallback(path, _w)


def write_txt(rows: List[Dict], path: Path) -> Path:
    def _w(p: Path) -> None:
        lines: List[str] = []
        done = [r for r in rows if r.get("status") in ("success", "partial")]
        failed = [r for r in rows if r.get("status") not in ("success", "partial")]
        for r in done:
            lines += ["=" * 50,
                      f"HR: {r.get('hr_name')}  ({r.get('hr_name_source')})",
                      f"EMAIL: {r.get('hr_email')}",
                      f"COMPANY: {r.get('company_name')}",
                      f"STATUS: {r.get('status')}" + (f"  - {r.get('notes')}" if r.get("notes") else ""),
                      "=" * 50, "",
                      "SUBJECT:", str(r.get("email_subject", "")), "",
                      "EMAIL:", str(r.get("email_body", "")), "",
                      "-" * 50, ""]
        if failed:
            lines += ["=" * 50, f"NOT GENERATED ({len(failed)})", "=" * 50]
            lines += [f"- {r.get('hr_email')}: {r.get('notes')}" for r in failed]
        p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return _write_with_fallback(path, _w)


def export(rows: List[Dict], output_dir: Path, fmt: str, also_txt: bool) -> List[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    written = []
    if fmt == "xlsx":
        written.append(write_xlsx(rows, output_dir / f"{BASE_NAME}.xlsx"))
    else:
        written.append(write_csv(rows, output_dir / f"{BASE_NAME}.csv"))
    if also_txt:
        written.append(write_txt(rows, output_dir / f"{BASE_NAME}.txt"))
    return written
